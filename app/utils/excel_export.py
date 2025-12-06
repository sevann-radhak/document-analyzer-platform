"""Excel export utility for exporting events to Excel format."""
import warnings
import io
from typing import List

# Suppress openpyxl deprecation warnings before importing
# These warnings come from openpyxl library using deprecated datetime.utcnow()
warnings.filterwarnings("ignore", message=".*datetime.datetime.utcnow.*", category=DeprecationWarning)
warnings.filterwarnings("ignore", category=DeprecationWarning, module="openpyxl")

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models.event import Event
from app.core.constants import ExcelExportConfig


def export_events_to_excel(events: List[Event]) -> io.BytesIO:
    """
    Export a list of events to Excel format.
    
    Creates an Excel workbook with formatted headers and event data.
    Includes styling for headers (blue background, white text) and borders.
    
    Args:
        events: List of Event objects to export
    
    Returns:
        BytesIO object containing the Excel file data
    
    Raises:
        ValueError: If events list is empty or invalid
        Exception: If Excel generation fails
    """
    if not isinstance(events, list):
        raise ValueError("Events must be a list")
    
    if not events:
        raise ValueError("Events list cannot be empty")
    
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = ExcelExportConfig.WORKSHEET_NAME
    
    _setup_header_row(worksheet)
    _populate_event_data(worksheet, events)
    _apply_column_widths(worksheet)
    
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    
    return output


def _setup_header_row(worksheet):
    """
    Set up the header row with styling.
    
    Args:
        worksheet: OpenPyXL worksheet object
    """
    headers = [
        ExcelExportConfig.HEADER_EVENT_ID,
        ExcelExportConfig.HEADER_EVENT_TYPE,
        ExcelExportConfig.HEADER_DESCRIPTION,
        ExcelExportConfig.HEADER_USER_ID,
        ExcelExportConfig.HEADER_CREATED_AT,
        ExcelExportConfig.HEADER_UPDATED_AT
    ]
    
    header_fill = PatternFill(
        start_color=ExcelExportConfig.HEADER_BACKGROUND_COLOR,
        end_color=ExcelExportConfig.HEADER_BACKGROUND_COLOR,
        fill_type="solid"
    )
    header_font = Font(
        bold=True,
        color=ExcelExportConfig.HEADER_TEXT_COLOR,
        size=ExcelExportConfig.HEADER_FONT_SIZE
    )
    header_alignment = Alignment(
        horizontal=ExcelExportConfig.ALIGNMENT_HORIZONTAL_CENTER,
        vertical=ExcelExportConfig.ALIGNMENT_VERTICAL_CENTER
    )
    border = Border(
        left=Side(style=ExcelExportConfig.BORDER_STYLE),
        right=Side(style=ExcelExportConfig.BORDER_STYLE),
        top=Side(style=ExcelExportConfig.BORDER_STYLE),
        bottom=Side(style=ExcelExportConfig.BORDER_STYLE)
    )
    
    for col_idx, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border


def _populate_event_data(worksheet, events: List[Event]):
    """
    Populate worksheet with event data.
    
    Args:
        worksheet: OpenPyXL worksheet object
        events: List of Event objects
    """
    border = Border(
        left=Side(style=ExcelExportConfig.BORDER_STYLE),
        right=Side(style=ExcelExportConfig.BORDER_STYLE),
        top=Side(style=ExcelExportConfig.BORDER_STYLE),
        bottom=Side(style=ExcelExportConfig.BORDER_STYLE)
    )
    
    for row_idx, event in enumerate(events, start=2):
        worksheet.cell(row=row_idx, column=1, value=event.id).border = border
        worksheet.cell(row=row_idx, column=2, value=event.event_type).border = border
        worksheet.cell(row=row_idx, column=3, value=event.description).border = border
        worksheet.cell(row=row_idx, column=4, value=event.user_id).border = border
        
        created_at_str = (
            event.created_at.strftime(ExcelExportConfig.DATE_FORMAT)
            if event.created_at else ""
        )
        worksheet.cell(row=row_idx, column=5, value=created_at_str).border = border
        
        updated_at_str = (
            event.updated_at.strftime(ExcelExportConfig.DATE_FORMAT)
            if event.updated_at else ""
        )
        worksheet.cell(row=row_idx, column=6, value=updated_at_str).border = border
        
        for col_idx in range(1, 7):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(
                vertical=ExcelExportConfig.ALIGNMENT_VERTICAL_TOP,
                wrap_text=True
            )


def _apply_column_widths(worksheet):
    """
    Apply appropriate column widths to the worksheet.
    
    Args:
        worksheet: OpenPyXL worksheet object
    """
    column_widths = {
        "A": ExcelExportConfig.COLUMN_WIDTH_EVENT_ID,
        "B": ExcelExportConfig.COLUMN_WIDTH_EVENT_TYPE,
        "C": ExcelExportConfig.COLUMN_WIDTH_DESCRIPTION,
        "D": ExcelExportConfig.COLUMN_WIDTH_USER_ID,
        "E": ExcelExportConfig.COLUMN_WIDTH_CREATED_AT,
        "F": ExcelExportConfig.COLUMN_WIDTH_UPDATED_AT
    }
    
    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

